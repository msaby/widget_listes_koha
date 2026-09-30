"""Préparer les ressources JSON et Excel des listes Koha (Python 3.10+)."""

# Ce fichier est le point d'entrée du traitement quotidien.
# Pour le lire dans l'ordre d'exécution : main() → convert() → read_documents(),
# puis CoverManager.populate(), write_excel() et enfin l'écriture du JSON.
# Un document est un dictionnaire Python : document["title"] contient son titre.
# La liste Python « documents » regroupe ces dictionnaires, une entrée par ligne CSV.

import argparse
from collections import defaultdict
from contextlib import contextmanager
import csv
from datetime import datetime, timezone
import json
import logging
from logging.handlers import RotatingFileHandler
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# L'import relatif fonctionne lorsque les tests importent scripts.prepare_koha_data.
# L'autre forme permet aussi de lancer directement « python scripts/prepare_koha_data.py ».
if __package__:
    from .covers import CoverManager, SOURCES
else:
    from covers import CoverManager, SOURCES


LOGGER = logging.getLogger(__name__)
# __file__ désigne ce script ; parents[1] remonte à la racine du dépôt.
# Les chemins par défaut ne dépendent donc pas du dossier depuis lequel on lance Python.
ROOT = Path(__file__).resolve().parents[1]
DEFAULT_LOG_FILE = ROOT / "logs/prepare_koha_data.log"
# FIELDS décrit le contrat de données, HEADERS les libellés publics de l'Excel.
# L'ordre des valeurs écrites par write_excel() doit correspondre à HEADERS.
FIELDS = (
    "biblionumber", "document_type", "title", "subtitle", "authors", "isbn",
    "ean", "date", "publisher", "pages", "abstract", "cover_url", "holdings",
    "opac_suppressed", "list_id", "list", "list_sort",
)
PRIMO_PREFIX = "https://catalogue.bu.univ-cotedazur.fr/permalink/f/11hf8mm/33UCA_KOHA"
HEADERS = (
    "Lien catalogue", "Type de document", "Titre", "Auteur", "ISBN", "EAN",
    "Date", "Éditeur", "Pages", "Résumé", "Exemplaires", "Liste",
)


def parse_holdings(value, context):
    """Tolérer les champs absents ; rejeter les structures non conformes."""
    if not value.strip():
        return []
    try:
        # Le CSV ne contient que du texte. json.loads transforme ici la cellule
        # '[{"library":"MEDP"}]' en une vraie liste contenant un dictionnaire.
        holdings = json.loads(value)
        if not isinstance(holdings, list) or any(
            not isinstance(item, dict)
            or any(not isinstance(v, str) for v in item.values())
            for item in holdings
        ):
            raise ValueError("tableau d'objets à valeurs textuelles attendu")
        return holdings
    except (ValueError, TypeError) as exc:
        # Une cellule d'exemplaires abîmée ne doit pas faire perdre toute la liste.
        # Le contexte indique dans les journaux quelle ligne/notice doit être corrigée.
        LOGGER.warning("%s : holdings invalide (%s), utilisation de []", context, exc)
        return []


def read_documents(path, delimiter=";"):
    # Retourne les documents en mémoire ; cette fonction n'écrit aucun fichier.
    documents = []
    # Les résumés peuvent dépasser la limite CSV par défaut de Python.
    csv.field_size_limit(10 * 1024 * 1024)
    # utf-8-sig accepte aussi le BOM, petit marqueur parfois ajouté par les exports.
    # newline="" laisse le module csv gérer les retours à la ligne dans les cellules.
    with Path(path).open(encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source, delimiter=delimiter, strict=True)
        headers = reader.fieldnames or []
        if len(set(headers)) != len(headers):
            raise ValueError("En-tête CSV : colonnes dupliquées")
        missing = set(FIELDS) - set(headers)
        # Une différence d'ensembles permet d'identifier les noms manquants/inconnus,
        # indépendamment de l'ordre des colonnes dans le fichier reçu.
        unknown = set(headers) - set(FIELDS)
        if missing or unknown:
            raise ValueError(
                f"En-tête CSV non conforme : manquantes={sorted(missing)}, "
                f"inconnues={sorted(unknown)}"
            )
        for row in reader:
            # DictReader associe chaque cellule au nom de sa colonne.
            # Une clé None ou une valeur None signale ici un nombre de cellules incorrect.
            context = f"Ligne {reader.line_num}, notice {row.get('biblionumber', '?')}"
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{context} : nombre de cellules incorrect")
            documents.append(normalize_document(row, context))
    return documents


def normalize_document(row, context):
    """Appliquer le contrat commun aux lignes des exports CSV et JSON Koha."""
    document = {}
    for field in FIELDS:
        value = row[field]
        if value is None:
            value = ""
        if isinstance(value, bool) or not isinstance(value, (str, int)):
            raise ValueError(f"{context} : {field} doit être du texte, un nombre ou null")
        document[field] = str(value)
    for field in ("biblionumber", "list_id"):
        value = document[field].strip()
        if not re.fullmatch(r"[0-9]+", value) or int(value) <= 0:
            raise ValueError(f"{context} : {field} doit être un entier positif")
        document[field] = str(int(value))
    document["holdings"] = parse_holdings(document["holdings"], context)
    document["record_url"] = PRIMO_PREFIX + document["biblionumber"]
    document["local_cover_url"] = ""
    return document


def read_json_documents(path):
    """Lire le JSON brut Koha : un tableau de lignes dans l'ordre de FIELDS."""
    with Path(path).open(encoding="utf-8-sig") as source:
        rows = json.load(source)
    if not isinstance(rows, list):
        raise ValueError("JSON Koha : tableau de lignes attendu")
    documents = []
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, list) or len(row) != len(FIELDS):
            raise ValueError(f"Ligne JSON {index} : {len(FIELDS)} colonnes attendues")
        documents.append(normalize_document(dict(zip(FIELDS, row)), f"Ligne JSON {index}, notice {row[0]}"))
    return documents


def format_holdings(holdings):
    # Prépare UNE cellule Excel : chaque exemplaire devient une ligne de texte.
    # join() ajoute les séparateurs entre les morceaux présents, jamais aux extrémités.
    lines = []
    for item in holdings:
        parts = [item.get(key, "") for key in ("library", "location")]
        if item.get("callnumber"):
            parts.append("cote " + item["callnumber"])
        line = " — ".join(part for part in parts if part)
        if line:
            lines.append(line)
    return "\n".join(lines)


@contextmanager
def atomic_destination(path):
    """Écrire dans le même dossier, puis remplacer seulement après succès."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        # Avec « with atomic_destination(...) as temporary », le code appelant
        # écrit dans temporary pendant la suspension de cette fonction au mot yield.
        # Si ce code lève une exception, os.replace n'est pas exécuté : l'ancien
        # fichier reste intact. Le finally nettoie le temporaire dans les deux cas.
        yield temporary
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_excel(documents, path):
    # openpyxl crée un classeur en mémoire. L'en-tête occupe la ligne 1 ; les
    # documents commencent donc à la ligne 2 (Excel numérote depuis 1, pas depuis 0).
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Documents"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="28556B")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    try:
        for row_number, document in enumerate(documents, start=2):
            title = " : ".join(filter(None, (document["title"], document["subtitle"])))
            values = (
                "Voir dans le catalogue", document["document_type"], title,
                document["authors"], document["isbn"], document["ean"],
                document["date"], document["publisher"], document["pages"],
                document["abstract"], format_holdings(document["holdings"]), document["list"],
            )
            for column, value in enumerate(values, start=1):
                # Excel a une limite par cellule. On refuse l'export de cette liste
                # plutôt que de publier silencieusement un résumé incomplet.
                if len(value) > 32767:
                    raise ValueError(
                        f"Notice {document['biblionumber']}, {HEADERS[column - 1]} : "
                        "limite Excel de 32767 caractères dépassée"
                    )
                cell = sheet.cell(row_number, column, value)
                # Les données bibliographiques sont du texte, jamais des formules.
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if column in (5, 6):
                    # Le format @ est le format « texte » d'Excel pour ISBN/EAN.
                    cell.number_format = "@"
            sheet.cell(row_number, 1).hyperlink = document["record_url"]
            sheet.cell(row_number, 1).style = "Hyperlink"
        widths = (25, 20, 55, 35, 22, 22, 12, 30, 22, 70, 48, 30)
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
        # Figer à A2 laisse l'en-tête visible pendant le défilement des lignes.
        # Le filtre porte sur tout le tableau, y compris les en-têtes.
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        with atomic_destination(path) as temporary:
            workbook.save(temporary)
    finally:
        workbook.close()


def convert(csv_path, json_path, exports_dir, delimiter=";", *, covers_dir=None,
            cover_sources=SOURCES, cover_timeout=10, cover_delay=0.3,
            google_api_key=None, retry_missing_covers=False):
    # Le * dans la signature impose de nommer les options suivantes lors de l'appel.
    # Cette fonction coordonne les étapes et retourne également le JSON sous forme
    # de dictionnaire Python, ce qui permet aux tests de vérifier son contenu.
    source_path = Path(csv_path)
    if source_path.suffix.lower() == ".json":
        source_format = "json"
        documents = read_json_documents(source_path)
    elif source_path.suffix.lower() == ".csv":
        source_format = "csv"
        documents = read_documents(source_path, delimiter)
    else:
        raise ValueError("--input doit désigner un fichier .csv ou .json")
    lists = {}
    grouped = defaultdict(list)
    # « lists » contient le nom et les réglages de chaque liste ; « grouped »
    # contient ses documents. defaultdict(list) crée automatiquement une liste
    # Python vide lorsqu'un list_id est rencontré pour la première fois.
    for document in documents:
        list_id = document["list_id"]
        metadata = {field: document[field] for field in ("list_id", "list", "list_sort")}
        if list_id in lists and lists[list_id] != metadata:
            raise ValueError(f"Métadonnées contradictoires pour la liste {list_id}")
        lists[list_id] = metadata
        grouped[list_id].append(document)
    if covers_dir is not None:
        # populate() enrichit les dictionnaires existants avec local_cover_url.
        # Les mêmes dictionnaires sont référencés dans documents et grouped.
        manager = CoverManager(covers_dir, cover_sources, timeout=cover_timeout,
                               delay=cover_delay, google_api_key=google_api_key,
                               retry_missing=retry_missing_covers)
        manager.populate(documents, json_path)
    exports_count = 0
    for list_id, selection in grouped.items():
        destination = Path(exports_dir) / f"liste-{list_id}.xlsx"
        try:
            # relpath calcule le chemin depuis le JSON, as_posix remplace les
            # séparateurs Windows par / et quote encode notamment les espaces.
            relative_url = quote(Path(os.path.relpath(
                destination.resolve(), Path(json_path).resolve().parent
            )).as_posix(), safe="/.")
            write_excel(selection, destination)
        except Exception as exc:
            LOGGER.error("Liste %s : échec de l'export Excel : %s", list_id, exc)
        else:
            # Le bloc else d'un try ne s'exécute que si aucune exception n'a été levée.
            # On ne publie donc un lien de téléchargement qu'après création du XLSX.
            lists[list_id]["export_xlsx"] = relative_url
            exports_count += 1
    payload = {
        "metadata": {
            "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "source": source_format,
            "document_count": len(documents),
            "list_count": len(lists),
        },
        "lists": lists,
        "documents": documents,
    }
    # ensure_ascii=False garde les accents lisibles ; indent=2 facilite l'inspection.
    # Cet export utilise tous les documents : recherche et pagination appartiennent
    # au navigateur et n'ont aucun effet sur les fichiers produits ici.
    with atomic_destination(json_path) as temporary:
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    LOGGER.info(
        "%d lignes lues, %d documents, %d listes, %d XLSX produits",
        len(documents), len(documents), len(lists), exports_count,
    )
    return payload


def configure_logging(path, verbose=False):
    """Conserver les avertissements et erreurs, et les infos sur demande."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    level = logging.INFO if verbose else logging.WARNING
    formatter = logging.Formatter("%(asctime)s %(levelname)s [%(name)s] %(message)s")
    file_handler = RotatingFileHandler(path, maxBytes=5 * 1024 * 1024,
                                       backupCount=3, encoding="utf-8")
    file_handler.setFormatter(formatter)
    console_handler = logging.StreamHandler()
    console_handler.setFormatter(formatter)
    logging.basicConfig(level=level, handlers=[file_handler, console_handler], force=True)
    return file_handler, console_handler


def main(argv=None):
    # argparse transforme les options du terminal en attributs de args et fournit
    # automatiquement --help. argv=None lit le terminal ; les tests passent une liste.
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "data/kohadata.csv")
    parser.add_argument("--output", type=Path, default=ROOT / "dist/prepareddata.json")
    parser.add_argument("--exports-dir", type=Path, default=ROOT / "exports")
    parser.add_argument("--delimiter", default=";")
    parser.add_argument("--covers-dir", type=Path, default=ROOT / "covers")
    parser.add_argument("--cover-source", choices=SOURCES, action="append",
                        help="Répéter pour définir l'ordre (défaut : google, bnf, amazon)")
    parser.add_argument("--cover-timeout", type=float, default=10)
    parser.add_argument("--cover-delay", type=float, default=0.3)
    parser.add_argument("--skip-covers", action="store_true", help="Conversion sans traitement des couvertures")
    parser.add_argument("--retry-missing-covers", action="store_true", help="Réinitialiser le cache négatif des fournisseurs")
    parser.add_argument("--verbose", action="store_true", help="Inclure les messages INFO dans le journal et le terminal")
    parser.add_argument("--log-file", type=Path, default=DEFAULT_LOG_FILE,
                        help="Journal de préparation (défaut : logs/prepare_koha_data.log)")
    args = parser.parse_args(argv)
    try:
        handlers = configure_logging(args.log_file, args.verbose)
    except OSError as exc:
        parser.exit(1, f"Impossible d'ouvrir le journal {args.log_file} : {exc}\n")
    try:
        if len(args.delimiter) != 1:
            raise ValueError("--delimiter doit contenir un seul caractère")
        payload = convert(args.input, args.output, args.exports_dir, args.delimiter,
                          covers_dir=None if args.skip_covers else args.covers_dir,
                          cover_sources=args.cover_source if args.cover_source is not None else SOURCES,
                          cover_timeout=args.cover_timeout, cover_delay=args.cover_delay,
                          google_api_key=os.environ.get("GOOGLE_BOOKS_API_KEY"),
                          retry_missing_covers=args.retry_missing_covers)
        # Le code de sortie est lu par le terminal ou le planificateur quotidien :
        # 0 = succès ; 1 = erreur de conversion ou au moins un Excel non produit.
        # Les échecs de couvertures sont journalisés, mais ne rendent pas ce code négatif.
        return 0 if all("export_xlsx" in item for item in payload["lists"].values()) else 1
    except (OSError, ValueError, csv.Error) as exc:
        LOGGER.error("Conversion interrompue : %s", exc)
        return 1
    except Exception:
        LOGGER.exception("Erreur imprévue pendant la préparation")
        return 1
    finally:
        # main() est aussi appelé depuis les tests ; libérer le journal sous Windows.
        for handler in handlers:
            logging.getLogger().removeHandler(handler)
            handler.close()


if __name__ == "__main__":
    # Ce bloc est exécuté au lancement direct, mais pas quand les tests importent le module.
    raise SystemExit(main())
