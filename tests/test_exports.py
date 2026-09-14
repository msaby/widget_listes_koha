# Tests des classeurs : on écrit un vrai XLSX, puis on le rouvre avec openpyxl
# pour inspecter ses cellules, ses liens et sa mise en forme. Aucun Excel installé
# ni aucune interaction manuelle n'est nécessaire pour ces vérifications.
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from scripts.csv_to_json import FIELDS, PRIMO_PREFIX, format_holdings, write_excel


class ExportTests(unittest.TestCase):
    def setUp(self):
        # Une notice de référence est recréée avant chaque test. Chaque scénario
        # peut donc la modifier sans influencer les suivants.
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.path = Path(self.temporary.name) / "liste-11.xlsx"
        self.document = dict.fromkeys(FIELDS, "") | {
            "biblionumber": "1", "list_id": "11", "list": "Fonds pro",
            "title": "Un titre", "subtitle": "Sous-titre", "isbn": "0012345",
            "ean": "0001234567890", "abstract": "Résumé intégral\n" * 1000,
            "record_url": PRIMO_PREFIX + "1", "holdings": [
                {"library": "MEDP", "location": "BUR", "callnumber": "Z 692 JAC"},
                {"library": "LASH", "callnumber": "É / 025.04"},
            ],
        }

    def test_workbook_content_and_formatting(self):
        # 40 documents + une ligne d'en-tête = 41 lignes attendues.
        # L'export doit être complet même au-delà de la limite par défaut du widget.
        write_excel([self.document] * 40, self.path)
        workbook = load_workbook(BytesIO(self.path.read_bytes()))
        # Lire depuis BytesIO évite de conserver un verrou Windows sur le fichier
        # temporaire. addCleanup ferme tout de même le classeur à la fin du test.
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 41)
        self.assertEqual([cell.value for cell in sheet[1]], [
            "Lien catalogue", "Type de document", "Titre", "Auteur", "ISBN", "EAN",
            "Date", "Éditeur", "Pages", "Résumé", "Exemplaires", "Liste",
        ])
        self.assertEqual(sheet["A2"].value, "Voir dans le catalogue")
        self.assertEqual(sheet["A2"].hyperlink.target, PRIMO_PREFIX + "1")
        self.assertEqual(sheet["C2"].value, "Un titre : Sous-titre")
        self.assertEqual(sheet["E2"].value, "0012345")
        self.assertEqual(sheet["F2"].value, "0001234567890")
        self.assertEqual(sheet["E2"].data_type, "s")
        # « s » désigne une chaîne de caractères ; « @ » le format texte d'Excel.
        # Vérifier la valeur ET le type empêche qu'un identifiant devienne un nombre.
        self.assertEqual(sheet["E2"].number_format, "@")
        self.assertEqual(sheet["J2"].value, self.document["abstract"])
        self.assertEqual(sheet["K2"].value, "MEDP — BUR — cote Z 692 JAC\nLASH — cote É / 025.04")
        for coordinate in ("J2", "K2"):
            self.assertTrue(sheet[coordinate].alignment.wrap_text)
            self.assertEqual(sheet[coordinate].alignment.vertical, "top")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:L41")
        self.assertTrue(sheet["A1"].font.bold)

    def test_absent_values_and_formula_like_text(self):
        # Un titre commençant par = doit être affiché tel quel, pas exécuté comme
        # une formule. Les cellules vides sont relues sous forme de None par openpyxl.
        self.document.update(title="=1+1", subtitle="", abstract="", holdings=[])
        write_excel([self.document], self.path)
        workbook = load_workbook(BytesIO(self.path.read_bytes()))
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet["C2"].value, "=1+1")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertIsNone(sheet["J2"].value)
        self.assertIsNone(sheet["K2"].value)

    def test_missing_holding_properties(self):
        # Les dictionnaires incomplets ne doivent pas produire de séparateurs isolés.
        self.assertEqual(format_holdings([]), "")
        self.assertEqual(format_holdings([{}, {"callnumber": "A"}, {"location": "BUR"}]), "cote A\nBUR")

    def test_excel_limit_does_not_silently_truncate_or_replace(self):
        # Placer un fichier témoin avant l'appel permet de prouver qu'une cellule
        # trop longue ne remplace pas un ancien classeur par un résultat partiel.
        self.path.write_bytes(b"previous workbook")
        self.document["abstract"] = "a" * 32768
        with self.assertRaisesRegex(ValueError, "32767"):
            write_excel([self.document], self.path)
        self.assertEqual(self.path.read_bytes(), b"previous workbook")


if __name__ == "__main__":
    unittest.main()
